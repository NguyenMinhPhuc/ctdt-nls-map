import openpyxl
import docx
import json
import re
import os

def build_data():
    print("Parsing Map_MonHoc_NLS_CapNhat.xlsx...")
    wb_capnhat = openpyxl.load_workbook('Map_MonHoc_NLS_CapNhat.xlsx', data_only=True)
    
    # 1. Tra cứu Khung NLS
    ws_nls = wb_capnhat['Tra Cứu Khung NLS (TT 02)']
    domains = {
        "I": "Khai thác dữ liệu và thông tin",
        "II": "Giao tiếp và hợp tác trong môi trường số",
        "III": "Sáng tạo nội dung số",
        "IV": "An toàn số",
        "V": "Giải quyết vấn đề bằng công nghệ số",
        "VI": "Ứng dụng Trí tuệ nhân tạo (AI)"
    }
    
    competencies = []
    # Rows 5 to 28 contain the 24 NLTP
    for r in range(5, 29):
        d_code = str(ws_nls.cell(r, 1).value or '').strip()
        d_name = str(ws_nls.cell(r, 2).value or '').strip()
        c_code = str(ws_nls.cell(r, 3).value or '').strip()
        c_name = str(ws_nls.cell(r, 4).value or '').strip()
        c_desc = str(ws_nls.cell(r, 5).value or '').strip()
        if c_code:
            competencies.append({
                "domainCode": d_code,
                "domainName": d_name or domains.get(d_code, ""),
                "code": c_code,
                "name": c_name,
                "desc": c_desc
            })
            
    # Levels description (rows 33 to 38)
    levels = []
    for r in range(33, 39):
        cap_do = str(ws_nls.cell(r, 1).value or '').strip()
        muc_do = str(ws_nls.cell(r, 2).value or '').strip()
        ten_muc_do = str(ws_nls.cell(r, 3).value or '').strip()
        mo_ta = str(ws_nls.cell(r, 4).value or '').strip()
        if muc_do:
            levels.append({
                "capDo": cap_do,
                "mucDo": muc_do,
                "tenMucDo": ten_muc_do,
                "moTa": mo_ta
            })

    # 2. Ma trận 6 Miền
    ws_m6 = wb_capnhat['Ma Trận Môn Học - 6 Miền NLS']
    matrix_6 = []
    for r in range(5, ws_m6.max_row+1):
        stt = ws_m6.cell(r, 1).value
        ma_mh = ws_m6.cell(r, 2).value
        ten_mh = ws_m6.cell(r, 3).value
        if not ma_mh or not ten_mh:
            continue
        matrix_6.append({
            "stt": stt,
            "maMH": str(ma_mh).strip(),
            "tenMH": str(ten_mh).strip(),
            "soTC": ws_m6.cell(r, 4).value,
            "khoiKT": str(ws_m6.cell(r, 5).value or '').strip(),
            "m1": str(ws_m6.cell(r, 6).value or '-').strip(),
            "m2": str(ws_m6.cell(r, 7).value or '-').strip(),
            "m3": str(ws_m6.cell(r, 8).value or '-').strip(),
            "m4": str(ws_m6.cell(r, 9).value or '-').strip(),
            "m5": str(ws_m6.cell(r, 10).value or '-').strip(),
            "m6": str(ws_m6.cell(r, 11).value or '-').strip(),
            "cacMien": str(ws_m6.cell(r, 12).value or '').strip(),
            "mucDoMax": str(ws_m6.cell(r, 13).value or '').strip()
        })

    # 3. Ma trận 24 NLTP
    ws_m24 = wb_capnhat['Ma Trận Môn Học - 24 NLTP']
    nltp_codes = [str(ws_m24.cell(5, c).value).strip() for c in range(6, 30)]
    matrix_24 = []
    for r in range(6, ws_m24.max_row+1):
        ma_mh = ws_m24.cell(r, 2).value
        ten_mh = ws_m24.cell(r, 3).value
        if not ma_mh or not ten_mh:
            continue
        m_map = {}
        for idx, code in enumerate(nltp_codes, start=6):
            val = ws_m24.cell(r, idx).value
            m_map[code] = True if val and str(val).strip().upper() == 'X' else False
        matrix_24.append({
            "maMH": str(ma_mh).strip(),
            "tenMH": str(ten_mh).strip(),
            "soTC": ws_m24.cell(r, 4).value,
            "khoiKT": str(ws_m24.cell(r, 5).value or '').strip(),
            "mapping": m_map
        })

    # 4. Gợi ý & Minh chứng
    ws_ev = wb_capnhat['Gợi Ý Bổ Sung & Minh Chứng NLS']
    evidences = []
    for r in range(5, ws_ev.max_row+1):
        ma_hp = ws_ev.cell(r, 2).value
        ten_hp = ws_ev.cell(r, 3).value
        if not ma_hp or not ten_hp:
            continue
        evidences.append({
            "stt": ws_ev.cell(r, 1).value,
            "maHP": str(ma_hp).strip(),
            "tenHP": str(ten_hp).strip(),
            "khoiKT": str(ws_ev.cell(r, 4).value or '').strip(),
            "mienCanBoSung": str(ws_ev.cell(r, 5).value or '').strip(),
            "maNLTP": str(ws_ev.cell(r, 6).value or '').strip(),
            "noiDungCanBoSung": str(ws_ev.cell(r, 7).value or '').strip(),
            "hoatDongDayHoc": str(ws_ev.cell(r, 8).value or '').strip(),
            "cachLayMinhChung": str(ws_ev.cell(r, 9).value or '').strip(),
            "mucDoHuongDen": str(ws_ev.cell(r, 10).value or '').strip()
        })

    # 5. Ánh xạ chi tiết CLO từ Template_Map.xlsx
    print("Parsing Template_Map.xlsx...")
    wb_tmpl = openpyxl.load_workbook('Template_Map.xlsx', data_only=True)
    ws_clo = wb_tmpl['NLS - CNTT']
    clo_mappings = []
    for r in range(18, ws_clo.max_row+1):
        ten_mh = ws_clo.cell(r, 2).value
        ma_mh = ws_clo.cell(r, 3).value
        if not ten_mh or not ma_mh:
            continue
        clo_mappings.append({
            "stt": ws_clo.cell(r, 1).value,
            "tenMH": str(ten_mh).strip(),
            "maMH": str(ma_mh).strip(),
            "soTC": ws_clo.cell(r, 4).value,
            "khoiKT": str(ws_clo.cell(r, 5).value or '').strip(),
            "mienNLS": str(ws_clo.cell(r, 6).value or '').strip(),
            "maNLTP": str(ws_clo.cell(r, 7).value or '').strip(),
            "tenNLTP": str(ws_clo.cell(r, 8).value or '').strip(),
            "mucDoNLS": str(ws_clo.cell(r, 9).value or '').strip(),
            "clo": str(ws_clo.cell(r, 10).value or '').strip(),
            "noiDungCLO": str(ws_clo.cell(r, 11).value or '').strip(),
            "pi": str(ws_clo.cell(r, 12).value or '').strip(),
            "danhGia": str(ws_clo.cell(r, 13).value or '').strip(),
            "ghiChu": str(ws_clo.cell(r, 14).value or '').strip()
        })

    # 6. Tóm tắt học phần từ docx
    print("Parsing 2026-CNTT-Tóm tắt học phần.docx...")
    doc = docx.Document('2026-CNTT-Tóm tắt học phần.docx')
    course_summaries = {}
    current_code = None
    current_title = None
    current_desc = []
    current_refs = []
    in_refs = False

    pattern = re.compile(r'^(\d{6})\s*[-–]\s*(.+)$')
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        m = pattern.match(txt)
        if m:
            if current_code:
                course_summaries[current_code] = {
                    "code": current_code,
                    "title": current_title,
                    "description": "\n\n".join(current_desc),
                    "references": "\n".join(current_refs)
                }
            current_code = m.group(1).strip()
            current_title = m.group(2).strip()
            current_desc = []
            current_refs = []
            in_refs = False
        elif current_code:
            if 'Tài liệu tham khảo' in txt:
                in_refs = True
            elif in_refs:
                current_refs.append(txt)
            else:
                current_desc.append(txt)

    if current_code:
        course_summaries[current_code] = {
            "code": current_code,
            "title": current_title,
            "description": "\n\n".join(current_desc),
            "references": "\n".join(current_refs)
        }

    # 7. Kế hoạch đào tạo từ BM03
    print("Parsing BM03-Ke hoach dao tao_CNTT-2026_guiDaotao.xlsx...")
    wb_bm03 = openpyxl.load_workbook('BM03-Ke hoach dao tao_CNTT-2026_guiDaotao.xlsx', data_only=True)
    ws_bm03 = wb_bm03['CTDT_2026']
    course_plan = {}
    cur_semester = "Chưa phân loại"
    
    for r in range(9, 140):
        val0 = str(ws_bm03.cell(r, 1).value or '').strip()
        val1 = str(ws_bm03.cell(r, 2).value or '').strip()
        
        if 'HỌC KỲ' in val0.upper() or 'HỌC KỲ' in val1.upper():
            cur_semester = val0 if 'HỌC KỲ' in val0.upper() else val1
            continue
            
        try:
            code_int = int(float(val0))
            code = str(code_int)
            course_plan[code] = {
                "code": code,
                "name": val1,
                "credits": ws_bm03.cell(r, 3).value,
                "totalHours": ws_bm03.cell(r, 4).value,
                "theoryHours": ws_bm03.cell(r, 5).value,
                "practiceHours": ws_bm03.cell(r, 6).value,
                "exerciseHours": ws_bm03.cell(r, 7).value,
                "note": str(ws_bm03.cell(r, 8).value or '').strip(),
                "semester": cur_semester
            }
        except:
            pass

    # Enrich Matrix 6 with Course Plan and Docx Info
    for c in matrix_6:
        code = c["maMH"]
        if code in course_plan:
            c["semester"] = course_plan[code]["semester"]
            c["theoryHours"] = course_plan[code]["theoryHours"]
            c["practiceHours"] = course_plan[code]["practiceHours"]
        else:
            c["semester"] = "N/A"
            c["theoryHours"] = None
            c["practiceHours"] = None
            
        if code in course_summaries:
            c["hasSummary"] = True
            c["summary"] = course_summaries[code]["description"]
            c["references"] = course_summaries[code]["references"]
        else:
            c["hasSummary"] = False
            c["summary"] = ""
            c["references"] = ""

    # Generate complete dataset
    unified_data = {
        "metadata": {
            "university": "TRƯỜNG ĐẠI HỌC LẠC HỒNG",
            "faculty": "KHOA CÔNG NGHỆ THÔNG TIN",
            "program": "CHƯƠNG TRÌNH ĐÀO TẠO CÔNG NGHỆ THÔNG TIN KHÓA 2026",
            "framework": "Thông tư 02/2025/TT-BGDĐT & Khung trình độ quốc gia VQF",
            "dateUpdated": "2026-09-13",
            "totalCourses": len(matrix_6),
            "totalCLOMappings": len(clo_mappings),
            "totalDomains": 6,
            "totalCompetencies": len(competencies),
            "totalEvidences": len(evidences)
        },
        "domains": domains,
        "competencies": competencies,
        "levels": levels,
        "matrix6": matrix_6,
        "matrix24": matrix_24,
        "nltpCodes": nltp_codes,
        "evidences": evidences,
        "cloMappings": clo_mappings,
        "courseSummaries": course_summaries,
        "coursePlan": course_plan
    }

    with open('curriculum_data.json', 'w', encoding='utf-8') as f:
        json.dump(unified_data, f, ensure_ascii=False, indent=2)
        
    print("Done! Saved curriculum_data.json successfully.")

if __name__ == "__main__":
    build_data()
